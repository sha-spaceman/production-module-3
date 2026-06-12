import streamlit as st
import pandas as pd
import sqlite3
import io
from openpyxl.styles import Font, PatternFill, Alignment

# ── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="QP/MS Parts Extractor",
    page_icon="📄",
    layout="centered",
)

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    /* Main background */
    .stApp { background-color: #f5f7fa; }

    /* Title bar accent */
    .title-block {
        background: linear-gradient(90deg, #2E5090 0%, #4472C4 100%);
        padding: 1.4rem 2rem 1rem 2rem;
        border-radius: 10px;
        margin-bottom: 1.5rem;
    }
    .title-block h1 { color: #ffffff !important; margin: 0; font-size: 1.8rem; }
    .title-block p  { color: #cdd9f0 !important; margin: 0.3rem 0 0 0; font-size: 0.95rem; }

    /* Cards */
    .card {
        background: #ffffff;
        border-radius: 10px;
        padding: 1.4rem 1.6rem;
        margin-bottom: 1.2rem;
        box-shadow: 0 2px 8px rgba(0,0,0,0.07);
    }
    .card h3 { margin-top: 0; color: #2E5090; font-size: 1.05rem; }

    /* Stat boxes */
    .stat-row { display: flex; gap: 1rem; margin-top: 0.5rem; }
    .stat-box {
        flex: 1;
        background: #eef2fb;
        border-left: 4px solid #2E5090;
        border-radius: 6px;
        padding: 0.7rem 1rem;
        font-size: 0.9rem;
        color: #2E5090;
    }
    .stat-box span { display: block; font-size: 1.6rem; font-weight: 700; color: #2E5090; }

    /* Radio label tweak */
    div[data-testid="stRadio"] label { font-size: 0.97rem; }

    /* Download button */
    .stDownloadButton > button {
        background-color: #2E5090 !important;
        color: white !important;
        border-radius: 6px !important;
        padding: 0.55rem 1.4rem !important;
        font-size: 0.97rem !important;
        width: 100%;
    }
    .stDownloadButton > button:hover { background-color: #1d3a6e !important; }

    /* Tab styling */
    .stTabs [data-baseweb="tab"] { font-size: 0.95rem; font-weight: 600; }
    .stTabs [aria-selected="true"] { color: #2E5090 !important; }
</style>
""", unsafe_allow_html=True)

# ── Parsers ───────────────────────────────────────────────────────────────────
def parse_caie(name):
    qp_ms = name[:14]
    q_num = name[15:17]
    q_part = name[18] if len(name) == 23 else "1"
    return qp_ms, q_num, q_part


def parse_edexcel(name):
    if name[7:8] == "r":
        qp_ms = name[:21]
        q_num = name[22:24]
    else:
        qp_ms = name[:20]
        q_num = name[21:23]
    if len(name) == 29:
        q_part = name[24]
    elif len(name) == 30:
        q_part = name[25]
    else:
        q_part = "1"
    return qp_ms, q_num, q_part


def parse_ibdp(name):
    qp_ms = name[:29]
    q_num = name[30:32]
    q_part = name[33] if len(name) == 38 else "1"
    return qp_ms, q_num, q_part


# ── Core logic ────────────────────────────────────────────────────────────────
def build_parts_table(filenames, board):
    parsers = {1: parse_caie, 2: parse_edexcel, 3: parse_ibdp}
    parser = parsers[board]
    rows = []
    for name in filenames:
        name = str(name).strip()
        if not name:
            continue
        qp_ms, q_num, q_part = parser(name)
        rows.append({
            "Q paper/Mark Scheme": qp_ms,
            "Q Number": q_num,
            "Q part": q_part,
        })
    return pd.DataFrame(rows, columns=["Q paper/Mark Scheme", "Q Number", "Q part"])


def run_queries(df, board):
    con = sqlite3.connect(":memory:")
    df.to_sql("File Parts info", con, index=False, if_exists="replace")

    if board in (1, 3):
        qp_filter, ms_filter = "%qp%", "%ms%"
    else:
        qp_filter, ms_filter = "%que%", "%rms%"

    sql = """
        SELECT
            "Q paper/Mark Scheme",
            "Q Number",
            MAX("Q part") AS "MaxOfQ part"
        FROM "File Parts info"
        GROUP BY "Q paper/Mark Scheme", "Q Number"
        HAVING "Q paper/Mark Scheme" LIKE ?
    """
    qp_df = pd.read_sql_query(sql, con, params=(qp_filter,))
    ms_df = pd.read_sql_query(sql, con, params=(ms_filter,))
    con.close()
    return qp_df, ms_df


def build_excel_bytes(qp_df, ms_df):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        qp_df.to_excel(writer, sheet_name="QParts", index=False)
        ms_df.to_excel(writer, sheet_name="MSParts", index=False)

        for sheet_name in ("QParts", "MSParts"):
            ws = writer.sheets[sheet_name]
            hfont = Font(name="Arial", bold=True, color="FFFFFF")
            hfill = PatternFill("solid", start_color="2E5090")
            for cell in ws[1]:
                cell.font = hfont
                cell.fill = hfill
                cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                max_len = max(len(str(c.value)) if c.value else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = max_len + 4
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial")
    buf.seek(0)
    return buf.read()


# ── UI ────────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="title-block">
  <h1>📄 QP / MS Parts Extractor</h1>
  <p>Examinent · Automated question paper and mark scheme parts analysis</p>
</div>
""", unsafe_allow_html=True)

# Step 1 – Board selection
st.markdown('<div class="card"><h3>① Select Examination Board</h3>', unsafe_allow_html=True)
board_label = st.radio(
    label="board",
    options=["Cambridge (CAIE)", "Pearson Edexcel", "IBDP"],
    label_visibility="collapsed",
)
board = {"Cambridge (CAIE)": 1, "Pearson Edexcel": 2, "IBDP": 3}[board_label]
st.markdown('</div>', unsafe_allow_html=True)

# Step 2 – File upload
st.markdown('<div class="card"><h3>② Upload PNG Filenames List (.xlsx)</h3>', unsafe_allow_html=True)
uploaded = st.file_uploader(
    label="Excel file with a single column of PNG filenames (no header)",
    type=["xlsx", "xls"],
    label_visibility="collapsed",
)
st.markdown('</div>', unsafe_allow_html=True)

# Step 3 – Process
if uploaded:
    raw = pd.read_excel(uploaded, header=None)
    filenames = raw.iloc[:, 0].dropna().tolist()

    if not filenames:
        st.error("No filenames found in the uploaded file. Please check the file contents.")
        st.stop()

    with st.spinner("Processing filenames…"):
        parts_df = build_parts_table(filenames, board)
        qp_df, ms_df = run_queries(parts_df, board)
        excel_bytes = build_excel_bytes(qp_df, ms_df)

    # Summary stats
    st.markdown(f"""
    <div class="card">
      <h3>③ Results Summary</h3>
      <div class="stat-row">
        <div class="stat-box">Total filenames<span>{len(filenames)}</span></div>
        <div class="stat-box">QParts rows<span>{len(qp_df)}</span></div>
        <div class="stat-box">MSParts rows<span>{len(ms_df)}</span></div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # Preview tables
    st.markdown('<div class="card"><h3>④ Preview</h3>', unsafe_allow_html=True)
    tab1, tab2, tab3 = st.tabs(["📋 File Parts Info", "📘 QParts", "📗 MSParts"])

    with tab1:
        st.dataframe(parts_df, use_container_width=True, height=260)

    with tab2:
        if qp_df.empty:
            st.info("No question paper records found for the selected board.")
        else:
            st.dataframe(qp_df, use_container_width=True, height=260)

    with tab3:
        if ms_df.empty:
            st.info("No mark scheme records found for the selected board.")
        else:
            st.dataframe(ms_df, use_container_width=True, height=260)

    st.markdown('</div>', unsafe_allow_html=True)

    # Download
    st.markdown('<div class="card"><h3>⑤ Download Output</h3>', unsafe_allow_html=True)
    st.download_button(
        label="⬇️  Download QP_MS_Parts.xlsx",
        data=excel_bytes,
        file_name="QP_MS_Parts.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.markdown('</div>', unsafe_allow_html=True)

else:
    st.info("Upload an Excel file above to begin processing.")

# ── Footer ────────────────────────────────────────────────────────────────────
st.markdown("---")
st.markdown(
    "<p style='text-align:center; color:#888; font-size:0.82rem;'>Examinent · QP/MS Parts Extractor</p>",
    unsafe_allow_html=True,
)
